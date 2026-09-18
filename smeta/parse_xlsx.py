# -*- coding: utf-8 -*-
"""Парсер локальных смет из xlsx в нормализованную таблицу.

Устроен как классификатор строк с последующей сборкой блоков, а не как «прочитать
таблицу с N-й строки». В настоящей выгрузке ГРАНД-Сметы одна позиция занимает
десяток строк:

    35 | 1 | ФЕР46-04-012-03 | Разборка заполнений | 100 м2 | 1,5387 |   <- сама позиция
    36 |   |                 | ОТ                  |        |          |   <- оплата труда
    37 |   |                 | ЭМ                  |        |          |   <- эксплуатация машин
    41 |   |                 | Итого по расценке   |        |          |
    43 |   | Приказ 812/пр   | НР Работы по реконс.| %      | 91       |   <- накладные
    44 |   | Приказ 774/пр   | СП Работы по реконс.| %      | 52       |   <- сметная прибыль
    45 |   |                 | Всего по позиции    |        |          |

Поэтому: сначала размечаем колонки по строке нумерации («1 2 3 4…» под шапкой —
она есть в любой выгрузке и переживает объединённые ячейки), потом классифицируем
каждую строку, потом собираем блоки в одну запись на позицию.
"""
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from .schema import (COLUMNS, KIND_OVERHEAD, KIND_POSITION, KIND_PROFIT,
                     KIND_SECTION, KIND_SUBTOTAL, KIND_TOTAL, KIND_UNKNOWN)

# --- разметка колонок -------------------------------------------------------

HEADER_ALIASES = {
    "num": ["№ пп", "№ п/п", "n пп", "номер"],
    "code": ["обоснование", "шифр", "код расценки"],
    "name": ["наименование работ", "наименование"],
    "unit": ["ед. изм", "ед.изм", "единица измерения", "ед изм"],
    "qty": ["количество", "кол-во", "объем", "объём"],
    "price": ["цена за ед", "цена на ед", "стоимость единицы", "сметная стоимость"],
    "coef": ["коэффициент", "к-т", "поправ"],
    "total": ["всего, руб", "общая стоимость", "всего", "сумма"],
    "index": ["индекс"],
    "ozp": ["озп", "осн. з/п", "основная заработная плата", "оплата труда"],
    "emm": ["эмм", "эксплуатация машин", "экспл. маш"],
    "mat": ["материалы", "мат."],
}

def sub_field(label):
    """Подзаголовок группы -> (роль, уровень цен).

    Порядок проверок важен: «всего с учётом коэффициентов» — это «всего», а не
    «коэффициенты». Уровень цен берётся из самого подзаголовка: в ресурсно-индексном
    методе базисная и текущая цена лежат в соседних колонках одной группы.
    """
    level = "текущий" if "текущ" in label else ("базисный" if "базисн" in label else None)
    if label.startswith("индекс"):
        return "index", level
    if label.startswith("всего"):
        return "total", level
    if label.startswith("на еди") or label.startswith("на ед"):
        return "unit", level
    if "коэффициент" in label:
        return "coef", level
    return None, level

SECTION_RE = re.compile(r"^\s*раздел\s+([\d.]+)?\.?\s*(.+?)\s*$", re.I)
SUBTOTAL_RE = re.compile(r"^\s*(итог\w*|всего|справочно|в том числе|строительные работы|"
                         r"монтажные работы|оплата труда|материалы|оборудование)\b", re.I)
TOTAL_RE = re.compile(r"^\s*(всего по смете|итого по смете|всего по локальн)", re.I)
PCT_RE = re.compile(r"(\d+[.,]?\d*)\s*%")
COEF_RE = re.compile(r"(?:к(?:оэф\w*)?[\s.]*=?\s*)(\d+[.,]\d+|\d+)", re.I)
CODE_RE = re.compile(
    r"^(ФЕР|ТЕР|ГЭСН|ФССЦ|ТССЦ|ФСБЦ|ФСЭМ|ТСЭМ|ФЕРм|ТЕРм|ФЕРр|ТЕРр|ГЭСНр|ГЭСНм|ГЭСНп)"
    # ФСБЦ-2022 нумерует ресурсы числами: 14.4.01.02-0012, машины 91.06.06-048,
    # трудозатраты 1-100-20. Это такая же нормативная база, как ФЕР с буквами
    r"|^\d{2}\.\d\.\d{2}\.\d{2}-\d{3,4}"
    r"|^\d{2}\.\d{2}\.\d{2}-\d{3,4}"
    r"|^\d{1,2}-\d{3}-\d{2,3}",
    re.I)
# Обоснование ссылкой на приказ: цена подтверждается конъюнктурным анализом
ORDER_REF_RE = re.compile(r"^(421/пр|Пр/\d+|Приказ)", re.I)

# строки внутри блока позиции: расшифровка стоимости, не отдельные работы
SUBROW_LABELS = {
    "от": "ozp", "оплата труда": "ozp", "эм": "emm", "эксплуатация машин": "emm",
    "в т.ч. отм": "ozpm", "в т.ч. от м": "ozpm", "отм(зтм)": "ozpm", "отм": "ozpm",
    "зт": "zt", "зтм": "ztm", "зт(зтм)": "zt",
    "фот": "fot", "мат": "mat", "м": "mat", "мр": "mat",
    "итого по расценке": "rate_total", "итого прямые затраты": "rate_total",
    "прямые затраты": "rate_total", "всего по позиции": "position_total",
}
# «Объём: 1,9*1,5*8+1,2*1,5» — расчёт объёма прямо в смете. Ценная строка:
# показывает, откуда взялся объём, и её стоит показывать в протоколе
VOLUME_CALC_RE = re.compile(r"^\s*об[ъь][её]м\s*:", re.I)
OVERHEAD_RE = re.compile(r"^\s*нр\b|наклад\w*\s+расход", re.I)
PROFIT_RE = re.compile(r"^\s*сп\b|сметн\w*\s+прибыл", re.I)
# «ОЗП=0,8; ЭМ=0,8 к расх.; МАТ=0 к расх.» или «ПЗ=6» — отдельная строка применения
# коэффициентов к позиции, со ссылкой на приказ или техническую часть сборника
# Без re.I и только отдельными словами: одиночная «М» матчилась внутри обычных
# наименований («…трубопроводов: 50»), и настоящие позиции уезжали в коэффициенты
COEF_RULE_RE = re.compile(
    r"(?<![А-Яа-яЁёA-Za-z])(ОЗП|ЭМм|ЭМ|ЗПМ|МАТ|ТЗМ|ТЗ|ПЗ|ЗТм|ЗТ)\s*[=:]\s*[\d.,]+"
    r"|^\s*[Цц]ена\s*="
    r"|^\s*[Фф]ормула\s+ценообразования")
COEF_SOURCE_RE = re.compile(r"^(приказ|тч|мдс|прил\w*|техническ)", re.I)


def to_number(value):
    """Число из ячейки: float, int или строка «1 234,56» / «1 234.56»."""
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    s = str(value).strip()
    if not s:
        return None
    s = (s.replace(" ", "").replace(" ", "").replace(" ", "")
          .replace("руб.", "").replace("руб", ""))
    # «100 м2» — единица измерения, а не число: без этой проверки остаток от
    # выброшенных букв склеивается в 1002 и уезжает в объём или цену
    if re.search(r"[A-Za-zА-Яа-яЁё]", s):
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    else:
        s = s.replace(",", ".")
    s = re.sub(r"[^\d.\-]", "", s)
    if s in ("", "-", ".", "-.") or s.count(".") > 1:
        return None
    try:
        return float(s)
    except ValueError:
        return None


# Номер позиции бывает с пометкой: «2 О» — оборудование, «5 М» — монтаж.
# Обычный разбор числа такую ячейку отвергает, и позиция приклеивается к предыдущей
POSITION_NUM_RE = re.compile(r"^\s*(\d+)\s*([А-Яа-яA-Za-z]{1,2})?\s*$")


def position_number(value):
    """Номер позиции по порядку. «2 О» -> 2.0, «Раздел» -> None."""
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    m = POSITION_NUM_RE.match(str(value).strip())
    return float(m.group(1)) if m else None


def parse_coef(text):
    """Произведение коэффициентов из ячейки. Пусто -> 1.0"""
    if text is None or str(text).strip() == "":
        return 1.0
    v = to_number(text)
    # ноль — законный коэффициент (МАТ=0 при демонтаже), и верхней границы почти нет:
    # из технической части сборников прилетают коэффициенты вида ПЗ=52
    if v is not None and 0 <= v <= 10000:
        return round(v, 6)
    result = 1.0
    for raw in COEF_RE.findall(str(text)):
        c = to_number(raw)
        if c and 0.05 < c < 20:
            result *= c
    return round(result, 6)


def _norm(s):
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def is_numbering_row(cells):
    """Строка «1 2 3 4 …» под шапкой — якорь разметки колонок."""
    vals = [(i, to_number(c)) for i, c in enumerate(cells) if c not in (None, "")]
    nums = [v for _, v in vals if v is not None and v == int(v)]
    if len(nums) < 5 or len(nums) != len(vals):
        return False
    ints = [int(v) for v in nums]
    return ints == list(range(ints[0], ints[0] + len(ints))) and ints[0] in (1, 0)


def is_header_row(cells):
    joined = " | ".join(_norm(c) for c in cells)
    return "наименование" in joined and ("обоснование" in joined or "количество" in joined)


def build_mapping(header_cells, sub_rows):
    """Шапка + подзаголовки -> поле: номер колонки.

    Верхняя строка задаёт группы («Количество», «Сметная стоимость в базисном
    уровне»), нижняя уточняет («на единицу», «коэффициенты», «всего»). Группа
    тянется от своей ячейки до следующей непустой — так объединённые ячейки
    восстанавливаются без чтения merged-диапазонов.
    """
    groups = []  # (col, текст, конец группы)
    marks = [(i, _norm(c)) for i, c in enumerate(header_cells) if _norm(c)]
    for k, (col, text) in enumerate(marks):
        end = marks[k + 1][0] if k + 1 < len(marks) else len(header_cells)
        groups.append((col, text, end))

    mapping = {}

    def claim(field, col):
        if field not in mapping and col is not None:
            mapping[field] = col

    for col, text, end in groups:
        field = None
        for name, aliases in HEADER_ALIASES.items():
            if any(a in text for a in aliases):
                field = name
                break
        if field is None:
            continue

        sub = {}
        for row in sub_rows:
            for i in range(col, min(end, len(row))):
                label = _norm(row[i])
                if label:
                    sub.setdefault(i, label)

        if field == "qty":
            if sub:
                for i, label in sub.items():
                    role, _ = sub_field(label)
                    if role == "unit":
                        claim("qty_unit", i)
                    elif role == "coef":
                        claim("qty_coef", i)
                    elif role == "total":
                        claim("qty", i)
                claim("qty", mapping.get("qty_unit"))
            else:
                claim("qty", col)
        elif field == "price":
            bare = re.sub(r"\(.*?\)", " ", text)  # уточнения в скобках не должны решать
            group_level = "текущий" if ("текущ" in bare and "базисн" not in bare) else "базисный"
            if sub:
                for i, label in sub.items():
                    role, level = sub_field(label)
                    level = level or group_level
                    if role == "index":
                        claim("index", i)
                    elif role == "unit":
                        claim("price_current" if level == "текущий" else "price", i)
                    elif role == "coef":
                        claim("price_coef", i)
                    elif role == "total":
                        claim("total_current" if level == "текущий" else "total", i)
            else:
                claim("total_current" if group_level == "текущий" else "price", col)
        else:
            claim(field, col)

    return mapping


def table_width(cells):
    last = -1
    for i, c in enumerate(cells):
        if c is not None and str(c).strip():
            last = i
    return last + 1


def find_table(rows, max_scan=120):
    """Первая строка данных, разметка колонок и ширина таблицы.

    Ширину берём по строке нумерации: правее неё ГРАНД-Смета печатает зеркальную
    копию тех же данных (колонки 75+), и если её не отрезать, она попадает и в
    разметку колонок, и в названия разделов.
    """
    for i, row in enumerate(rows[:max_scan]):
        if not is_header_row(row):
            continue
        sub_rows, data_start, width = [], i + 1, table_width(row)
        for j in range(i + 1, min(i + 6, len(rows))):
            if is_numbering_row(rows[j]):
                width = table_width(rows[j])
                data_start = j + 1
                break
            if any(_norm(c) for c in rows[j]):
                sub_rows.append(rows[j])
            data_start = j + 1
        mapping = build_mapping(row[:width], [r[:width] for r in sub_rows])
        if "name" in mapping and ("total" in mapping or "qty" in mapping):
            return data_start, mapping, width
    return None, {}, 0


def cell(cells, mapping, field):
    idx = mapping.get(field)
    if idx is None or idx >= len(cells):
        return None
    v = cells[idx]
    return v.strip() if isinstance(v, str) else v


def text_of(cells, mapping, field):
    v = cell(cells, mapping, field)
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


# --- классификация и сборка -------------------------------------------------

def classify(cells, mapping):
    if not any(c not in (None, "") for c in cells):
        return None, ""
    if is_header_row(cells) or is_numbering_row(cells):
        return None, ""

    name = text_of(cells, mapping, "name")
    code = text_of(cells, mapping, "code")
    joined = " ".join(str(c) for c in cells if isinstance(c, str) and c.strip())

    if re.match(r"^\s*раздел\s", joined, re.I):
        return KIND_SECTION, joined
    if OVERHEAD_RE.search(name) or "812/пр" in code:
        return KIND_OVERHEAD, name
    if PROFIT_RE.search(name) or "774/пр" in code:
        return KIND_PROFIT, name

    label = SUBROW_LABELS.get(_norm(name))
    if label:
        return "sub:" + label, name
    if VOLUME_CALC_RE.match(name):
        return "sub:volume_calc", name
    if COEF_RULE_RE.search(name) or (COEF_SOURCE_RE.match(code) and not to_number(cell(cells, mapping, "total"))):
        return "sub:coef_rule", (code + " " + name).strip()
    if TOTAL_RE.match(name):
        return KIND_TOTAL, name
    if SUBTOTAL_RE.match(name):
        return KIND_SUBTOTAL, name

    num = position_number(cell(cells, mapping, "num"))
    qty = to_number(cell(cells, mapping, "qty"))
    total = to_number(cell(cells, mapping, "total"))
    if name and (code or num is not None) and (qty is not None or total is not None):
        # Ресурс внутри позиции («Вода | м3 | 3,5 | 2,44») выглядит как позиция:
        # есть шифр, объём и цена. Отличает её отсутствие номера по порядку —
        # без этой проверки ресурс обрывает блок, и сама позиция остаётся без цены.
        if "num" in mapping and num is None:
            return "sub:resource", name
        return KIND_POSITION, name
    # Ресурс без номера по порядку: шифр и наименование есть, а объём может быть
    # не числом, а буквой «П» — принимается по проекту. Это всё равно ресурс.
    if code and name and "num" in mapping and num is None and not SUBTOTAL_RE.match(name):
        return "sub:resource", name

    if not name and not code and qty is None and total is None:
        # Подраздел: текст лежит в колонке номера или обоснования, всё остальное
        # пусто. Молча пропускать нельзя — без него материалы из разных операций
        # выглядят дублями внутри одного раздела.
        for field in ("num", "code"):
            head = text_of(cells, mapping, field)
            if (head and to_number(head) is None and position_number(head) is None
                    and len(head) > 3
                    and not CODE_RE.match(head)):
                return "subsection", head
        return None, ""
    return KIND_UNKNOWN, name


def blank_record(**kw):
    rec = dict.fromkeys(COLUMNS)
    rec.update(kw)
    return rec


def parse_sheet(ws, file_name):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    data_start, mapping, width = find_table(rows)
    if data_start is None:
        return []
    rows = [r[:width] for r in rows]

    records = []
    section = None
    subsection = None   # блок работ внутри раздела: «Установка арматуры…»
    current = None      # собираемая позиция
    summary_mode = False  # идём внутри итогового блока раздела или сметы

    def flush():
        nonlocal current
        if current is None:
            return
        pos = current["rec"]
        parts = current["parts"]
        # цена за единицу и итог: «Итого по расценке» точнее, чем сама строка позиции
        rate = parts.get("rate_total")
        if rate:
            pos["price"] = rate.get("price") or pos["price"]
            pos["total"] = rate.get("total") or pos["total"]
            pos["total_current"] = rate.get("total_current") or pos["total_current"]
            pos["price_current"] = rate.get("price_current") or pos["price_current"]
        whole = parts.get("position_total")
        if whole:
            if whole.get("total"):
                pos["total_with_overhead"] = whole["total"]
            pos["total_current"] = pos["total_current"] or whole.get("total_current")
            pos["price_current"] = pos["price_current"] or whole.get("price_current")
        # В ресурсно-индексном методе базисного итога может не быть вовсе:
        # тогда считаем в текущем уровне цен и помечаем это в записи
        if pos.get("total") is None and pos.get("total_current") is not None:
            pos["total"] = pos["total_current"]
            pos["price_level"] = "текущий"
        else:
            pos["price_level"] = "базисный"
        # Итог в текущем уровне цен требует и цены в текущем уровне: иначе
        # проверка арифметики сравнивает базисную цену с текущей суммой
        if pos.get("price_level") == "текущий":
            if pos.get("price_current") is not None:
                pos["price"] = pos["price_current"]
            elif pos.get("price") is not None and pos.get("index"):
                pos["price"] = round(pos["price"] * pos["index"], 4)
        elif pos.get("price") is None and pos.get("price_current") is not None:
            pos["price"] = pos["price_current"]
        if current["coefs"]:
            existing = pos.get("coef_text")
            pos["coef_text"] = "; ".join(([existing] if existing else []) + current["coefs"])[:500]
        for src, dst in (("ozp", "ozp"), ("emm", "emm"), ("mat", "mat")):
            part = parts.get(src)
            if part and part.get("total") is not None:
                pos[dst] = part["total"]
        if pos.get("mat") is None and current["resources"]:
            spent = [r["total"] for r in current["resources"] if r.get("total") is not None]
            if spent:
                pos["mat"] = round(sum(spent), 2)
        pos["resources"] = len(current["resources"])
        pos["components"] = [c for c in current["components"] if c["kind"] != "ozpm"]
        nr = [e for e in current["extras"] if e["kind"] == KIND_OVERHEAD]
        sp = [e for e in current["extras"] if e["kind"] == KIND_PROFIT]
        pos["nr_sum"] = nr[0]["total"] if nr else None
        pos["sp_sum"] = sp[0]["total"] if sp else None
        records.append(pos)
        records.extend(current["extras"])
        current = None

    for offset, cells in enumerate(rows[data_start:], start=data_start + 1):
        kind, name = classify(cells, mapping)
        if kind is None:
            continue
        if kind in (KIND_SUBTOTAL, KIND_TOTAL):
            summary_mode = True
        elif kind in (KIND_POSITION, KIND_SECTION):
            summary_mode = False
        elif summary_mode and kind == KIND_UNKNOWN:
            kind = KIND_SUBTOTAL  # расшифровка итогов раздела

        if kind == KIND_SECTION:
            flush()
            m = SECTION_RE.match(name.strip())
            section = (m.group(2) if m else name).strip()
            subsection = None
            continue

        if kind == "subsection":
            flush()
            subsection = name.strip()
            continue

        vals = dict(
            coef=parse_coef(cell(cells, mapping, "price_coef")),
            index=to_number(cell(cells, mapping, "index")),
            num=position_number(cell(cells, mapping, "num")),
            qty=to_number(cell(cells, mapping, "qty")),
            qty_unit=to_number(cell(cells, mapping, "qty_unit")),
            price=to_number(cell(cells, mapping, "price")),
            price_current=to_number(cell(cells, mapping, "price_current")),
            total=to_number(cell(cells, mapping, "total")),
            total_current=to_number(cell(cells, mapping, "total_current")),
            unit=text_of(cells, mapping, "unit") or None,
            code=text_of(cells, mapping, "code") or None,
        )
        coef_text = cell(cells, mapping, "price_coef")
        if coef_text is None:
            coef_text = cell(cells, mapping, "coef")
        if coef_text is None:
            coef_text = cell(cells, mapping, "qty_coef")

        if kind == KIND_POSITION:
            flush()
            rec = blank_record(
                file=file_name, sheet=ws.title, excel_row=offset, kind=KIND_POSITION,
                section=section, subsection=subsection, num=vals["num"], code=vals["code"], name=name,
                unit=vals["unit"], qty=vals["qty"], price=vals["price"],
                price_current=vals["price_current"], coef=parse_coef(coef_text), coef_text=str(coef_text) if coef_text else None,
                total=vals["total"], total_current=vals["total_current"],
                index=vals["index"],
            )
            current = dict(rec=rec, parts={}, extras=[], coefs=[], resources=[], components=[])
            continue

        if kind == "sub:coef_rule":
            if current is not None:
                current["coefs"].append(name)
            continue

        if kind == "sub:volume_calc":
            if current is not None and not current["rec"].get("volume_calc"):
                current["rec"]["volume_calc"] = name[:200]
            continue

        if kind == "sub:resource":
            if current is not None:
                current["resources"].append(dict(name=name, **vals))
            continue

        if kind.startswith("sub:"):
            if current is not None:
                part = kind[4:]
                current["parts"][part] = vals
                if part in ("ozp", "emm", "mat", "ozpm"):
                    # составляющие стоимости: у каждой свой коэффициент, и проверять
                    # арифметику надо по ним, а не по сумме — коэффициенты разные
                    current["components"].append(dict(kind=part, name=name, **vals))
            continue

        if kind in (KIND_OVERHEAD, KIND_PROFIT):
            pct = to_number(cell(cells, mapping, "qty"))
            if pct is None:
                m = PCT_RE.search(name)
                pct = to_number(m.group(1)) if m else None
            work_type = re.sub(r"^\s*(нр|сп)\s+", "", name, flags=re.I).strip()
            rec = blank_record(
                file=file_name, sheet=ws.title, excel_row=offset, kind=kind,
                section=section, subsection=subsection, code=vals["code"], name=name, pct=pct,
                total=vals["total"], total_current=vals["total_current"],
                work_type=work_type or None,
                parent_row=current["rec"]["excel_row"] if current else None,
            )
            if current is not None:
                current["extras"].append(rec)
                current["rec"]["nr_pct" if kind == KIND_OVERHEAD else "sp_pct"] = pct
            else:
                records.append(rec)
            continue

        flush()
        records.append(blank_record(
            file=file_name, sheet=ws.title, excel_row=offset, kind=kind, section=section,
            subsection=subsection, name=name, code=vals["code"], total=vals["total"],
            total_current=vals["total_current"], qty=vals["qty"], price=vals["price"],
        ))

    flush()
    return records


def parse_file(path):
    """Книга -> DataFrame по схеме COLUMNS. Каждый лист может быть отдельной сметой."""
    path = Path(path)
    wb = load_workbook(path, data_only=True, read_only=True)
    records = []
    try:
        for ws in wb.worksheets:
            records.extend(parse_sheet(ws, path.name))
    finally:
        wb.close()
    if not records:
        return pd.DataFrame(columns=COLUMNS)
    return pd.DataFrame(records, columns=COLUMNS)


def parse_dir(folder, pattern="*.xls*"):
    """Все сметы папки в одну таблицу + сводка качества разбора по листам.

    Единица учёта — лист, а не файл: в одной книге лежит десяток локальных смет,
    и «разобрано 1 из 1 файла» ничего не говорит о том, сколько смет прочитано.
    """
    folder = Path(folder)
    frames, stats = [], []
    for f in sorted(folder.glob(pattern)):
        if f.name.startswith("~$"):
            continue
        try:
            df = parse_file(f)
        except Exception as e:  # битый файл не должен ронять прогон по корпусу
            stats.append(dict(file=f.name, sheet="—", positions=0, overheads=0, sections=0,
                              unknown=0, no_price=0, clean=False,
                              error="{}: {}".format(type(e).__name__, str(e)[:60])))
            continue
        if df.empty:
            stats.append(dict(file=f.name, sheet="—", positions=0, overheads=0, sections=0,
                              unknown=0, no_price=0, clean=False, error="не локальная смета"))
            continue
        frames.append(df)
        for sheet, part in df.groupby("sheet", sort=False):
            positions = int((part["kind"] == KIND_POSITION).sum())
            unknown = int((part["kind"] == KIND_UNKNOWN).sum())
            stats.append(dict(
                file=f.name,
                sheet=sheet,
                positions=positions,
                overheads=int(part["kind"].isin([KIND_OVERHEAD, KIND_PROFIT]).sum()),
                sections=int(part["section"].nunique()),
                unknown=unknown,
                no_price=int(((part["kind"] == KIND_POSITION) & part["price"].isna()).sum()),
                clean=positions > 0 and unknown == 0,
                error="",
            ))
    all_df = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=COLUMNS)
    return all_df, pd.DataFrame(stats)
