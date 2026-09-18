# -*- coding: utf-8 -*-
"""Генератор тестовых локальных смет в формате, близком к выгрузке Гранд-Сметы в xlsx.

Нужен, пока нет доступа к корпусу с zakupki.gov.ru: даёт парсеру те же болячки,
что и реальная выгрузка (объединённые ячейки, повтор шапки на разрыве страницы,
многострочные наименования, числа-строки, подытоги вперемешку с позициями),
плюс заложенные дефекты с ground truth — по нему считается recall валидатора.
"""
import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "fixtures"
TRUTH = ROOT / "data" / "fixtures_truth.json"

# Нормативы НР/СП — демонстрационные ориентиры по видам работ.
# Перед показом сверить с приказами Минстроя 812/пр и 774/пр: там свои цифры и своя разбивка.
SECTIONS = {
    "Земляные работы": dict(nr=95, sp=50),
    "Бетонные и железобетонные конструкции": dict(nr=105, sp=64),
    "Кровельные работы": dict(nr=108, sp=55),
    "Отделочные работы": dict(nr=105, sp=55),
}

# шифр, наименование, ед. изм., базовая цена за ед., доли ОЗП / ЭММ / МАТ
CATALOG = {
    "Земляные работы": [
        ("ФЕР01-01-013-02", "Разработка грунта с погрузкой на автомобили-самосвалы\nэкскаваторами с ковшом вместимостью 0,65 м3, группа грунтов 2", "1000 м3", 41850.0, .12, .61, .27),
        ("ФЕР01-01-030-01", "Разработка грунта вручную в траншеях глубиной до 2 м\nбез креплений с откосами, группа грунтов 1", "100 м3", 78400.0, .93, .07, .00),
        ("ФЕР01-02-001-01", "Уплотнение грунта прицепными катками на пневмоколёсном ходу\n25 т, группа грунтов 2", "1000 м3", 12300.0, .09, .78, .13),
        ("ФЕР01-01-036-03", "Засыпка траншей и котлованов с перемещением грунта\nбульдозерами мощностью 96 кВт", "1000 м3", 9750.0, .05, .88, .07),
    ],
    "Бетонные и железобетонные конструкции": [
        ("ФЕР06-01-001-13", "Устройство бетонной подготовки под фундаменты", "100 м3", 412600.0, .31, .18, .51),
        ("ФЕР06-01-015-07", "Устройство железобетонных фундаментных плит\nплоских толщиной до 500 мм", "100 м3", 587300.0, .28, .16, .56),
        ("ФЕР06-01-041-02", "Устройство стен подвала железобетонных высотой до 3 м,\nтолщиной до 300 мм", "100 м3", 731500.0, .34, .19, .47),
        ("ФССЦ-204-0025", "Горячекатаная арматурная сталь класса А-III диаметром 12 мм", "т", 58900.0, .00, .00, 1.0),
    ],
    "Кровельные работы": [
        ("ФЕР12-01-002-09", "Устройство кровель плоских из наплавляемых материалов\nв два слоя", "100 м2", 63400.0, .41, .07, .52),
        ("ФЕР12-01-013-01", "Устройство пароизоляции обмазочной в один слой", "100 м2", 11850.0, .48, .05, .47),
        ("ФССЦ-101-1791", "Материал рулонный кровельный наплавляемый, верхний слой", "м2", 412.0, .00, .00, 1.0),
    ],
    "Отделочные работы": [
        ("ФЕР15-02-016-03", "Штукатурка поверхностей внутри здания цементно-известковым\nраствором по камню стен", "100 м2", 87200.0, .72, .04, .24),
        ("ФЕР15-04-005-04", "Окраска водно-дисперсионными акриловыми составами\nулучшенная по штукатурке стен", "100 м2", 24600.0, .69, .03, .28),
        ("ФЕР11-01-011-01", "Устройство покрытий из керамогранитных плит\nразмером 30х30 см", "100 м2", 143900.0, .58, .06, .36),
    ],
}

THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BODY_FONT = Font(size=9)

COLS = ["№ пп", "Обоснование", "Наименование работ и затрат", "Ед. изм.", "Количество",
        "Цена за ед., руб.", "Коэффициенты", "Всего, руб.", "ОЗП", "ЭММ", "Материалы"]
WIDTHS = [6, 20, 52, 10, 12, 15, 16, 16, 12, 12, 12]


def rub(x):
    return round(x + 0.0, 2)


def as_text_number(x):
    """Число, записанное текстом с неразрывными пробелами и запятой — классика выгрузок."""
    return f"{x:,.2f}".replace(",", " ").replace(".", ",")


class SmetaBuilder:
    def __init__(self, wb, title, obj_name, number):
        self.ws = wb.active
        self.ws.title = "Локальная смета"
        self.row = 1
        self.defects = []
        self._titles(title, obj_name, number)

    def _titles(self, title, obj_name, number):
        ws = self.ws
        for i, w in enumerate(WIDTHS, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

        def line(text, bold=False, size=11):
            ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=11)
            c = ws.cell(row=self.row, column=1, value=text)
            c.font = Font(bold=bold, size=size)
            c.alignment = Alignment(horizontal="center", wrap_text=True)
            self.row += 1

        line(obj_name, size=10)
        line("(наименование стройки)", size=8)
        self.row += 1
        line("ЛОКАЛЬНЫЙ СМЕТНЫЙ РАСЧЁТ № " + number, bold=True, size=12)
        line(title, bold=True, size=10)
        line("(наименование работ и затрат, наименование объекта)", size=8)
        self.row += 1
        line("Основание: рабочая документация, шифр 24-117-АР", size=9)
        line("Составлен(а) в текущих (прогнозных) ценах по состоянию на 2 кв. 2026 г.", size=9)
        self.row += 1

    def header(self):
        ws = self.ws
        for i, name in enumerate(COLS, start=1):
            c = ws.cell(row=self.row, column=i, value=name)
            c.font = Font(bold=True, size=9)
            c.border = BORDER
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.row += 1
        for i in range(1, 12):  # строка нумерации колонок — есть в реальной выгрузке
            c = ws.cell(row=self.row, column=i, value=i)
            c.font = Font(size=8, italic=True)
            c.border = BORDER
            c.alignment = Alignment(horizontal="center")
        self.row += 1

    def section(self, name):
        ws = self.ws
        ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=11)
        c = ws.cell(row=self.row, column=1, value="Раздел " + name)
        c.font = Font(bold=True, size=10)
        self.row += 1

    def position(self, num, code, name, unit, qty, price, coef_text, coef_value,
                 ozp_share, emm_share, mat_share, total_override=None, defect=None):
        ws = self.ws
        total = rub(qty * price * coef_value) if total_override is None else rub(total_override)
        vals = [num, code, name, unit, qty, price, coef_text, total,
                rub(total * ozp_share), rub(total * emm_share), rub(total * mat_share)]
        for i, v in enumerate(vals, start=1):
            # часть числовых ячеек пишем текстом — парсер обязан это переварить
            if i in (5, 6, 8) and isinstance(v, float) and random.random() < 0.18:
                v = as_text_number(v)
            c = ws.cell(row=self.row, column=i, value=v)
            c.font = BODY_FONT
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top",
                                    horizontal="left" if i == 3 else "center")
        pos_row = self.row
        self.row += 1
        if defect:
            self.defects.append(dict(row=pos_row, code=code, **defect))
        return pos_row, total, rub(total * ozp_share)

    def overhead(self, fot, nr_pct, sp_pct, defect=None):
        ws = self.ws
        for label, pct in (("Накладные расходы", nr_pct), ("Сметная прибыль", sp_pct)):
            tag = "НР" if label.startswith("Наклад") else "СП"
            ws.cell(row=self.row, column=2, value=tag).font = BODY_FONT
            c = ws.cell(row=self.row, column=3, value=f"{label} {pct}% от ФОТ")
            c.font = Font(size=9, italic=True)
            c.alignment = Alignment(wrap_text=True)
            ws.cell(row=self.row, column=8, value=rub(fot * pct / 100)).font = BODY_FONT
            if defect and tag == "НР":
                self.defects.append(dict(row=self.row, code="НР", **defect))
            self.row += 1

    def subtotal(self, label, value):
        ws = self.ws
        c = ws.cell(row=self.row, column=3, value=label)
        c.font = Font(bold=True, size=9)
        v = ws.cell(row=self.row, column=8, value=rub(value))
        v.font = Font(bold=True, size=9)
        self.row += 1

    def spacer(self, n=1):
        self.row += n


def build_smeta(path, number, title, obj_name, section_names, rng, planted):
    wb = Workbook()
    b = SmetaBuilder(wb, title, obj_name, number)
    b.header()

    num = 0
    grand = 0.0
    rows_seen = []
    page_counter = 0

    for idx, sec_name in enumerate(section_names, start=1):
        norms = SECTIONS[sec_name]
        b.section(f"{idx}. {sec_name}")
        sec_total = 0.0
        items = list(CATALOG[sec_name])
        rng.shuffle(items)
        for code, name, unit, base_price, ozp, emm, mat in items:
            num += 1
            qty = round(rng.uniform(1.2, 9.4), 2)
            price = rub(base_price * rng.uniform(0.94, 1.07))
            coef_text, coef_value = "", 1.0
            total_override = None
            defect = None

            plant = planted.get((sec_name, code))
            if plant == "price":
                price = rub(base_price * rng.uniform(1.55, 1.9))
                defect = dict(rule="price_above_median", note="цена завышена относительно корпуса")
            elif plant == "round_qty":
                qty = float(rng.choice([500, 1000, 250]))
                defect = dict(rule="round_volume", note="объём подозрительно круглый")
            elif plant == "arith":
                total_override = qty * price * rng.uniform(1.08, 1.14)
                defect = dict(rule="arithmetic", note="сумма не равна объём умножить на цену и коэффициенты")
            elif plant == "coef_on_mat":
                coef_text, coef_value = "К=1,15 (стеснённые условия)", 1.15
                defect = dict(rule="coef_on_materials", note="повышающий коэффициент на материальную позицию")
            elif plant == "no_basis":
                code = "Прайс поставщика"
                defect = dict(rule="market_price", note="позиция по прайсу без обоснования расценкой")

            if not coef_text and rng.random() < 0.25:
                coef_text, coef_value = "К=1,1 (зимнее удорожание)", 1.1

            _, total, fot = b.position(num, code, name, unit, qty, price, coef_text,
                                       coef_value, ozp, emm, mat, total_override, defect)
            rows_seen.append((code, name, unit, qty, price, total))
            sec_total += total

            if ozp > 0.05:
                nr_defect = None
                nr = norms["nr"]
                if planted.get((sec_name, code + "#nr")) == "nr":
                    nr = norms["nr"] + rng.randint(25, 45)
                    nr_defect = dict(rule="overhead_out_of_range",
                                     note="НР " + str(nr) + "% при нормативе " + str(norms["nr"]) + "% для вида работ")
                b.overhead(fot, nr, norms["sp"], nr_defect)
                sec_total += rub(fot * nr / 100) + rub(fot * norms["sp"] / 100)

            page_counter += 1
            if page_counter % 7 == 0:  # разрыв страницы: шапка печатается заново
                b.spacer()
                b.header()

        # дубль: та же работа второй раз в конце раздела
        if planted.get((sec_name, "#dup")):
            code, name, unit, qty, price, total = rows_seen[-2]
            num += 1
            b.position(num, code, name, unit, qty, price, "", 1.0, .5, .2, .3, total,
                       defect=dict(rule="duplicate", note="позиция уже посчитана в этой смете"))
            sec_total += total

        b.subtotal("Итого по разделу " + sec_name, sec_total)
        b.spacer()
        grand += sec_total

    b.subtotal("ВСЕГО по смете", grand)
    wb.save(path)
    return b.defects, grand


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    rng = random.Random(20260824)
    random.seed(20260824)

    plan = [
        ("smeta_01_zemlyanye.xlsx", "02-01-01", "Земляные работы и устройство котлована",
         ["Земляные работы", "Бетонные и железобетонные конструкции"],
         {("Земляные работы", "ФЕР01-01-013-02"): "price",
          ("Земляные работы", "#dup"): True,
          ("Бетонные и железобетонные конструкции", "ФССЦ-204-0025"): "coef_on_mat"}),
        ("smeta_02_karkas.xlsx", "02-02-01", "Каркас здания. Монолитные конструкции",
         ["Бетонные и железобетонные конструкции", "Кровельные работы"],
         {("Бетонные и железобетонные конструкции", "ФЕР06-01-015-07#nr"): "nr",
          ("Кровельные работы", "ФССЦ-101-1791"): "no_basis"}),
        ("smeta_03_krovlya.xlsx", "02-03-01", "Кровля и водоизоляционный ковёр",
         ["Кровельные работы", "Отделочные работы"],
         {("Кровельные работы", "ФЕР12-01-002-09"): "arith",
          ("Отделочные работы", "#dup"): True}),
        ("smeta_04_otdelka.xlsx", "02-04-01", "Отделочные работы. Секция 1",
         ["Отделочные работы", "Земляные работы"],
         {("Земляные работы", "ФЕР01-01-030-01"): "round_qty",
          ("Отделочные работы", "ФЕР15-02-016-03#nr"): "nr"}),
        ("smeta_05_blagoustroystvo.xlsx", "02-05-01", "Благоустройство территории",
         ["Земляные работы", "Отделочные работы"], {}),
        ("smeta_06_podzemnaya.xlsx", "02-06-01", "Подземная часть. Фундаментная плита",
         ["Бетонные и железобетонные конструкции", "Земляные работы"],
         {("Бетонные и железобетонные конструкции", "ФЕР06-01-041-02"): "arith"}),
    ]

    truth = []
    for fname, number, title, sections, planted in plan:
        defects, total = build_smeta(OUT / fname, number, title,
                                     "Жилой комплекс по адресу: г. Москва, ул. Примерная, вл. 12",
                                     sections, rng, planted)
        for d in defects:
            truth.append(dict(file=fname, **d))
        print(fname + ": разделов " + str(len(sections)) +
              ", дефектов заложено " + str(len(defects)) +
              ", итого " + f"{total:,.2f}".replace(",", " ") + " руб.")

    TRUTH.write_text(json.dumps(truth, ensure_ascii=False, indent=2), encoding="utf-8")
    print("\nВсего заложено дефектов: " + str(len(truth)) + " -> " + str(TRUTH.relative_to(ROOT)))


if __name__ == "__main__":
    main()
